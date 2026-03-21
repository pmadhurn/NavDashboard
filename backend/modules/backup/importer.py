import csv
import io
import json
import logging
import zipfile
from datetime import datetime
from uuid import UUID

import openpyxl
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from modules.auth.models import User
from modules.couples.models import Couple
from modules.devices.models import Device
from modules.inventory.models import FittingMaterial, MaterialTemplate
from modules.locations.models import Location, LocationHistory
from modules.pairs.models import Pair
from modules.personnel.models import Person
from modules.status.models import StatusChangeLog
from modules.troubleshooting.models import ErrorLog, TroubleshootEntry
from shared.audit import record_audit

logger = logging.getLogger(__name__)


SHEET_MODEL_MAP = {
    "Devices": {
        "model": Device,
        "lookup_fields": ["id", "serial_number"],
    },
    "Couples": {
        "model": Couple,
        "lookup_fields": ["id", "name"],
    },
    "Pairs": {
        "model": Pair,
        "lookup_fields": ["id", "name"],
    },
    "Locations": {
        "model": Location,
        "lookup_fields": ["id"],
    },
    "LocationHistory": {
        "model": LocationHistory,
        "lookup_fields": ["id"],
    },
    "Personnel": {
        "model": Person,
        "lookup_fields": ["id", "full_name"],
    },
    "FittingMaterials": {
        "model": FittingMaterial,
        "lookup_fields": ["id", "name"],
    },
    "MaterialTemplates": {
        "model": MaterialTemplate,
        "lookup_fields": ["id", "template_name"],
    },
    "ErrorLogs": {
        "model": ErrorLog,
        "lookup_fields": ["id"],
    },
    "TroubleshootEntries": {
        "model": TroubleshootEntry,
        "lookup_fields": ["id"],
    },
    "StatusChangeLogs": {
        "model": StatusChangeLog,
        "lookup_fields": ["id"],
    },
    "Users": {
        "model": User,
        "lookup_fields": ["id", "email", "username"],
    },
}

CSV_KEY_MAP = {
    "devices": "Devices",
    "couples": "Couples",
    "pairs": "Pairs",
    "locations": "Locations",
    "location_history": "LocationHistory",
    "personnel": "Personnel",
    "fitting_materials": "FittingMaterials",
    "material_templates": "MaterialTemplates",
    "error_logs": "ErrorLogs",
    "troubleshoot_entries": "TroubleshootEntries",
    "status_change_logs": "StatusChangeLogs",
    "users": "Users",
}


def _parse_value(value: str, col_name: str):
    if value is None or value == "" or value == "None":
        return None
    if col_name == "id" or col_name.endswith("_id") or col_name.endswith("_by"):
        try:
            return UUID(value)
        except (ValueError, AttributeError):
            return None
    if col_name.endswith("_at") or col_name in ("moved_at", "reported_at", "performed_at", "changed_at"):
        try:
            return datetime.fromisoformat(value)
        except (ValueError, TypeError):
            return None
    if col_name in ("is_active", "is_template", "resolved", "had_rf"):
        return value.lower() in ("true", "1", "yes")
    if col_name in ("quantity", "step_number"):
        try:
            return int(value)
        except (ValueError, TypeError):
            return 0
    if col_name in ("latitude", "longitude", "old_latitude", "old_longitude",
                     "new_latitude", "new_longitude", "distance_meters"):
        try:
            return float(value)
        except (ValueError, TypeError):
            return None
    if col_name in ("metadata_json", "custom_fields", "materials",
                     "fitting_materials_snapshot", "configuration_snapshot"):
        try:
            return json.loads(value)
        except (json.JSONDecodeError, TypeError):
            return None
    return value


SKIP_COLUMNS = {"hashed_password", "password", "coordinate", "deleted_at"}


async def _find_existing(db: AsyncSession, model, lookup_fields: list[str], row_data: dict):
    for field in lookup_fields:
        val = row_data.get(field)
        if val is not None:
            try:
                col = getattr(model, field)
                stmt = select(model).where(col == val)
                result = await db.execute(stmt)
                existing = result.scalar_one_or_none()
                if existing:
                    return existing
            except Exception:
                continue
    return None


async def _upsert_rows(
    db: AsyncSession,
    model,
    lookup_fields: list[str],
    headers: list[str],
    rows: list[list[str]],
    user_id: UUID,
) -> dict:
    created = 0
    updated = 0
    errors = 0
    error_details = []

    for row_idx, raw_row in enumerate(rows, 2):
        try:
            row_data = {}
            for col_idx, col_name in enumerate(headers):
                if col_name in SKIP_COLUMNS:
                    continue
                if col_idx < len(raw_row):
                    parsed = _parse_value(raw_row[col_idx], col_name)
                    if parsed is not None:
                        row_data[col_name] = parsed

            existing = await _find_existing(db, model, lookup_fields, row_data)

            if existing:
                update_data = {k: v for k, v in row_data.items() if k != "id"}
                for key, val in update_data.items():
                    if hasattr(existing, key):
                        setattr(existing, key, val)
                updated += 1
            else:
                new_instance = model(**row_data)
                db.add(new_instance)
                created += 1

            await db.flush()

        except Exception as e:
            errors += 1
            error_details.append(f"Row {row_idx}: {str(e)[:200]}")
            logger.warning(f"Import error at row {row_idx}: {e}")

    return {
        "created": created,
        "updated": updated,
        "errors": errors,
        "error_details": error_details if error_details else None,
    }


async def import_xlsx(db: AsyncSession, file_bytes: bytes, user_id: UUID) -> dict:
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True)

    total_created = 0
    total_updated = 0
    total_errors = 0
    total_rows = 0
    all_error_details = []

    for sheet_name in wb.sheetnames:
        mapping = SHEET_MODEL_MAP.get(sheet_name)
        if not mapping:
            continue

        ws = wb[sheet_name]
        rows_list = list(ws.iter_rows(values_only=True))
        if len(rows_list) < 2:
            continue

        headers = [str(h) if h else "" for h in rows_list[0]]
        data_rows = []
        for row in rows_list[1:]:
            data_rows.append([str(cell) if cell is not None else "" for cell in row])

        total_rows += len(data_rows)
        result = await _upsert_rows(
            db, mapping["model"], mapping["lookup_fields"],
            headers, data_rows, user_id,
        )

        total_created += result["created"]
        total_updated += result["updated"]
        total_errors += result["errors"]
        if result["error_details"]:
            all_error_details.extend(
                [f"[{sheet_name}] {e}" for e in result["error_details"]]
            )

    wb.close()

    return {
        "total_rows": total_rows,
        "created": total_created,
        "updated": total_updated,
        "errors": total_errors,
        "error_details": all_error_details if all_error_details else None,
    }


async def import_csv(db: AsyncSession, zip_bytes: bytes, user_id: UUID) -> dict:
    total_created = 0
    total_updated = 0
    total_errors = 0
    total_rows = 0
    all_error_details = []

    try:
        zf = zipfile.ZipFile(io.BytesIO(zip_bytes), "r")
    except zipfile.BadZipFile:
        csv_content = zip_bytes.decode("utf-8", errors="replace")
        reader = csv.reader(io.StringIO(csv_content))
        all_rows = list(reader)
        if len(all_rows) < 2:
            return {
                "total_rows": 0, "created": 0, "updated": 0,
                "errors": 1, "error_details": ["CSV file is empty or has no data rows"],
            }

        return {
            "total_rows": len(all_rows) - 1,
            "created": 0, "updated": 0, "errors": 1,
            "error_details": ["Single CSV import: could not determine target table. Use ZIP format."],
        }

    for csv_filename in zf.namelist():
        key = csv_filename.replace(".csv", "").strip("/")
        sheet_name = CSV_KEY_MAP.get(key)
        if not sheet_name:
            continue

        mapping = SHEET_MODEL_MAP.get(sheet_name)
        if not mapping:
            continue

        csv_content = zf.read(csv_filename).decode("utf-8", errors="replace")
        reader = csv.reader(io.StringIO(csv_content))
        all_rows = list(reader)

        if len(all_rows) < 2:
            continue

        headers = all_rows[0]
        data_rows = all_rows[1:]
        total_rows += len(data_rows)

        result = await _upsert_rows(
            db, mapping["model"], mapping["lookup_fields"],
            headers, data_rows, user_id,
        )

        total_created += result["created"]
        total_updated += result["updated"]
        total_errors += result["errors"]
        if result["error_details"]:
            all_error_details.extend(
                [f"[{sheet_name}] {e}" for e in result["error_details"]]
            )

    zf.close()

    return {
        "total_rows": total_rows,
        "created": total_created,
        "updated": total_updated,
        "errors": total_errors,
        "error_details": all_error_details if all_error_details else None,
    }